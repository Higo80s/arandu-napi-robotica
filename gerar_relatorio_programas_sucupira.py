"""
Relatório Excel: Programas de Pós-Graduação por Campus
Formato oficial UTFPR/Sucupira - 81 cursos
"""

from app import app, db
from models import ProgramaAcademico
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO
from datetime import datetime

def gerar_relatorio_programas_sucupira():
    """
    Gera relatório Excel com todos os programas de pós-graduação
    Formato: Código | Nome | Sigla | Campus | Nível | Área Aval | Área Básica | Ano | Nota
    """
    with app.app_context():
        # Buscar todos os programas, exceto graduação
        programas = ProgramaAcademico.query.filter(
            ProgramaAcademico.niveis != 'GRAD',
            ~ProgramaAcademico.codigo_programa.like('GRAD-%')
        ).order_by(
            ProgramaAcademico.campus_id,  # Ordenar por campus
            ProgramaAcademico.sigla_programa
        ).all()
        
        # Criar workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Programas Pos-Graduacao"
        
        # TÍTULO
        ws.merge_cells('A1:I1')
        titulo_cell = ws['A1']
        titulo_cell.value = f"Relação dos programas de pós-graduação por câmpus - Atualizada em {datetime.now().strftime('%d/%m/%Y')}"
        titulo_cell.font = Font(size=14, bold=True, color="FFFFFF")
        titulo_cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
        titulo_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # CABEÇALHOS
        headers = [
            'Código - PPG',
            'Nome do curso',
            'Sigla',
            'Campus/campi',
            'Nivel',
            'Área de Avaliação - CAPES',
            'Área Básica - CAPES',
            'Ano Inicio',
            'Nota CAPES'
        ]
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        # Definir larguras de coluna
        ws.column_dimensions['A'].width = 18  # Código
        ws.column_dimensions['B'].width = 50  # Nome
        ws.column_dimensions['C'].width = 12  # Sigla
        ws.column_dimensions['D'].width = 12  # Campus
        ws.column_dimensions['E'].width = 22  # Nível
        ws.column_dimensions['F'].width = 35  # Área Avaliação
        ws.column_dimensions['G'].width = 35  # Área Básica
        ws.column_dimensions['H'].width = 10  # Ano
        ws.column_dimensions['I'].width = 10  # Nota
        
        # DADOS
        row_num = 3
        
        # Mapear níveis para texto legível
        niveis_map = {
            'M': 'Mestrado Acadêmico',
            'MP': 'Mestrado Profissional',
            'D': 'Doutorado Acadêmico',
            'DP': 'Doutorado Profissional',
            'M,D': ['Mestrado Acadêmico', 'Doutorado Acadêmico'],
            'MP,DP': ['Mestrado Profissional', 'Doutorado Profissional'],
            'MP,D': ['Mestrado Profissional', 'Doutorado Acadêmico']
        }
        
        # Para programas com múltiplos níveis, criar linhas separadas
        for prog in programas:
            niveis_prog = prog.niveis or ''
            
            # Verificar se tem múltiplos níveis
            if niveis_prog in niveis_map and isinstance(niveis_map[niveis_prog], list):
                # Criar linha para cada nível
                for nivel_texto in niveis_map[niveis_prog]:
                    row_data = [
                        prog.codigo_programa or '',
                        prog.nome_programa or '',
                        prog.sigla_programa or '',
                        prog.campus.sigla_campus if prog.campus else '',
                        nivel_texto,
                        prog.area_avaliacao or '',
                        prog.area_basica or '',
                        prog.ano_inicio or '',
                        prog.nota_capes or ''
                    ]
                    
                    for col_num, value in enumerate(row_data, 1):
                        cell = ws.cell(row=row_num, column=col_num)
                        cell.value = value
                        cell.alignment = Alignment(horizontal='left' if col_num in [2] else 'center', vertical='center')
                        
                        # Bordas
                        thin_border = Border(
                            left=Side(style='thin'),
                            right=Side(style='thin'),
                            top=Side(style='thin'),
                            bottom=Side(style='thin')
                        )
                        cell.border = thin_border
                    
                    row_num += 1
            else:
                # Nível único
                nivel_texto = niveis_map.get(niveis_prog, niveis_prog)
                
                row_data = [
                    prog.codigo_programa or '',
                    prog.nome_programa or '',
                    prog.sigla_programa or '',
                    prog.campus.sigla_campus if prog.campus else '',
                    nivel_texto,
                    prog.area_avaliacao or '',
                    prog.area_basica or '',
                    prog.ano_inicio or '',
                    prog.nota_capes or ''
                ]
                
                for col_num, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_num, column=col_num)
                    cell.value = value
                    cell.alignment = Alignment(horizontal='left' if col_num in [2] else 'center', vertical='center')
                    
                    thin_border = Border(
                        left=Side(style='thin'),
                        right=Side(style='thin'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin')
                    )
                    cell.border = thin_border
                
                row_num += 1
        
        # Altura das linhas de cabeçalho
        ws.row_dimensions[2].height = 40
        
        # Salvar em BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output

if __name__ == "__main__":
    # Teste local
    output = gerar_relatorio_programas_sucupira()
    with open('teste_relatorio_programas.xlsx', 'wb') as f:
        f.write(output.getvalue())
    print("Relatorio gerado: teste_relatorio_programas.xlsx")
