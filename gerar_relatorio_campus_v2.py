"""
Relatório Consolidado por Campus - VERSÃO REFORMULADA

Lógica:
1. CAPES DS (Edital 14): Contar TODAS as cotas distribuídas por campus (da tabela distribuicao_cotas)
2. CAPES IES (Edital 13): Contar bolsistas ATIVOS
3. Outros editais: Contar bolsistas ATIVOS

Baseado no dashboard CAPES e requisitos do usuário
"""

from app import app, db
from models import Bolsista, ProgramaAcademico
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import datetime
from io import BytesIO
from collections import defaultdict
from sqlalchemy import text

def gerar_relatorio_campus_novo():
    """
    Gera relatório Excel consolidado com nova lógica:
    - CAPES DS (14): Todas as cotas distribuídas
    - CAPES IES (13): Bolsistas ativos
    - Outros: Bolsistas ativos
    """
    
    dados = defaultdict(lambda: {'M': 0, 'D': 0})
    
    # 1. CAPES DS - Contar TODAS as cotas da tabela distribuicao_cotas (Edital 14)
    print("Processando CAPES DS (Edital 14)...")
    
    result = db.session.execute(text("""
        SELECT 
            c.sigla_campus,
            m.nivel,
            SUM(dc.qtd_concedida) as total
        FROM distribuicao_cotas dc
        JOIN programa_academico p ON dc.programa_id = p.id
        JOIN campus c ON p.campus_id = c.id
        JOIN modalidades m ON dc.modalidade_id = m.id
        WHERE dc.edital_id = 14
        GROUP BY c.sigla_campus, m.nivel
    """))
    
    for row in result:
        campus_sigla = row[0]
        nivel_raw = row[1].upper()
        total = row[2]
        
        if 'MESTRADO' in nivel_raw:
            dados[campus_sigla]['M'] += total
        elif 'DOUTORADO' in nivel_raw:
            dados[campus_sigla]['D'] += total
    
    # 2. CAPES IES - Bolsistas ATIVOS (Edital 13)
    print("Processando CAPES IES (Edital 13)...")
    
    bolsistas_ies = Bolsista.query.filter_by(edital_id=13, ativo=True).all()
    
    for b in bolsistas_ies:
        # Determinar campus
        if b.programa_link and b.programa_link.campus:
            campus_sigla = b.programa_link.campus.sigla_campus
        elif b.programa and '(' in b.programa:
            campus_sigla = b.programa.split('(')[-1].strip(')')
        else:
            campus_sigla = 'NÃO INFORMADO'
        
        # Determinar nível
        if not b.modalidade:
            continue
            
        nivel_raw = b.modalidade.nivel.upper()
        if 'MESTRADO' in nivel_raw:
            dados[campus_sigla]['M'] += 1
        elif 'DOUTORADO' in nivel_raw:
            dados[campus_sigla]['D'] += 1
    
    # 3. OUTROS EDITAIS - Bolsistas ATIVOS (exceto 13 e 14)
    print("Processando outros editais...")
    
    outros_bolsistas = Bolsista.query.filter(
        Bolsista.ativo == True,
        Bolsista.edital_id.notin_([13, 14])
    ).all()
    
    for b in outros_bolsistas:
        # Determinar campus
        if b.programa_link and b.programa_link.campus:
            campus_sigla = b.programa_link.campus.sigla_campus
        elif b.programa and '(' in b.programa:
            campus_sigla = b.programa.split('(')[-1].strip(')')
        else:
            campus_sigla = 'NÃO INFORMADO'
        
        # Determinar nível
        if not b.modalidade:
            continue
            
        nivel_raw = b.modalidade.nivel.upper()
        if 'MESTRADO' in nivel_raw or 'MASTER' in nivel_raw:
            dados[campus_sigla]['M'] += 1
        elif 'DOUTORADO' in nivel_raw or 'DOCTOR' in nivel_raw:
            dados[campus_sigla]['D'] += 1
    
    # Criar Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Bolsas por Campus"
    
    # Estilos
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    total_font = Font(bold=True, size=11)
    
    # Título
    ws.merge_cells('A1:D1')
    ws['A1'] = "Demonstrativo de Bolsas de Pós-graduação - UTFPR"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # Subtítulo explicativo
    ws.merge_cells('A2:D2')
    ws['A2'] = "CAPES DS (cotas totais) + CAPES IES (ativos) + Outros Editais (ativos)"
    ws['A2'].font = Font(italic=True, size=10, color="666666")
    ws['A2'].alignment = Alignment(horizontal='center')
    
    # Cabeçalho
    row = 4
    headers = ['Campus', 'Mestrado', 'Doutorado', 'TOTAL']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    # Ordenar campus
    campus_ordenados = sorted(dados.keys())
    
    # Dados
    row = 5
    total_m = 0
    total_d = 0
    
    for campus in campus_ordenados:
        m = dados[campus]['M']
        d = dados[campus]['D']
        total = m + d
        
        ws.cell(row=row, column=1, value=campus)
        ws.cell(row=row, column=2, value=m)
        ws.cell(row=row, column=3, value=d)
        ws.cell(row=row, column=4, value=total)
        
        for col in [2, 3, 4]:
            ws.cell(row=row, column=col).alignment = Alignment(horizontal='center')
        
        total_m += m
        total_d += d
        row += 1
    
    # Totais
    ws.cell(row=row, column=1, value='TOTAL').font = total_font
    
    ws.cell(row=row, column=2, value=total_m)
    ws.cell(row=row, column=2).font = total_font
    ws.cell(row=row, column=2).fill = total_fill
    ws.cell(row=row, column=2).alignment = Alignment(horizontal='center')
    
    ws.cell(row=row, column=3, value=total_d)
    ws.cell(row=row, column=3).font = total_font
    ws.cell(row=row, column=3).fill = total_fill
    ws.cell(row=row, column=3).alignment = Alignment(horizontal='center')
    
    ws.cell(row=row, column=4, value=total_m + total_d)
    ws.cell(row=row, column=4).font = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(row=row, column=4).fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws.cell(row=row, column=4).alignment = Alignment(horizontal='center')
    
    # Rodapé
    row += 2
    ws.merge_cells(f'A{row}:D{row}')
    ws.cell(row=row, column=1, value=f"Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
    ws.cell(row=row, column=1).font = Font(italic=True, size=9)
    
    row += 1
    ws.merge_cells(f'A{row}:D{row}')
    ws.cell(row=row, column=1, value="Fonte: Sistema de Gestão de Editais PROPPG")
    ws.cell(row=row, column=1).font = Font(italic=True, size=9)
    
    # Ajustar larguras
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    
    # Salvar
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    print(f"\nTotais: {total_m}M + {total_d}D = {total_m + total_d}")
    
    return output

# Testar
if __name__ == "__main__":
    from app import app
    with app.app_context():
        gerar_relatorio_campus_novo()
        print("✅ Relatório gerado!")
